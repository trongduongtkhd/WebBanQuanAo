from pathlib import Path
from copy import deepcopy
import json, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent.parent/'src/Docs/BaoCaoDoAn3_BoSung_2_2_1.docx'
FONT='C:/Windows/Fonts/arial.ttf'
BOLD='C:/Windows/Fonts/arialbd.ttf'
F=ImageFont.truetype(FONT,27)
FB=ImageFont.truetype(BOLD,28)
FS=ImageFont.truetype(FONT,25)
W=1700
specs=[]
def m(a,b,s): return ('m',a,b,s)
def r(a,b,s): return ('r',a,b,s)
def alt(s): return ('start','alt',s)
def opt(s): return ('start','opt',s)
def loop(s): return ('start','loop',s)
def other(s): return ('else',s)
def end(): return ('end',)
def ref(s): return ('ref',s)
def add(title,req,pre,desc,events,service='Xử lý nghiệp vụ',actor='Khách hàng',lanes=None):
 specs.append(dict(title=title,req=req,pre=pre,desc=desc,events=events,lanes=lanes or [actor,'Giao diện\nAngular',service,'Dữ liệu\nEF Core / SQL']))

add('Luồng xử lý chung của hệ thống','R1.7, R15','Người dùng truy cập một chức năng của hệ thống.',
 'Giao diện gửi yêu cầu đến API. Pipeline kiểm tra xác thực và quyền đối với endpoint được bảo vệ; endpoint công khai không bắt buộc JWT. Controller tiếp nhận dữ liệu, Service thực hiện nghiệp vụ và truy cập dữ liệu qua AppDbContext. Controller tạo ApiResponse khi thành công; ExceptionMiddleware xử lý ngoại lệ. Lỗi xác thực hoặc phân quyền có thể kết thúc yêu cầu ngay trong pipeline.',
 [m(0,1,'Chọn chức năng, nhập dữ liệu'),m(1,2,'Gửi HTTP request'),m(2,2,'Xác thực và phân quyền\n(nếu endpoint yêu cầu)'),alt('Không đủ điều kiện truy cập'),r(2,1,'401 hoặc 403'),other('Được phép truy cập'),m(2,3,'Controller gọi Service'),m(3,4,'Truy vấn / cập nhật qua EF Core'),r(4,3,'Kết quả dữ liệu'),r(3,2,'DTO hoặc ngoại lệ'),r(2,1,'ApiResponse hoặc phản hồi lỗi'),end(),r(1,0,'Hiển thị kết quả')],lanes=['Người dùng','Angular','Pipeline API','Controller\n/ Service','AppDbContext\n/ SQL Server'])
add('Đăng ký tài khoản','R1.1','Người dùng chưa có tài khoản và mở biểu mẫu đăng ký.',
 'Sau khi kiểm tra dữ liệu và CAPTCHA, AuthService chuẩn hóa email, kiểm tra trùng và lấy vai trò Customer. Mật khẩu được băm bằng BCrypt trước khi lưu User và UserRole. Hệ thống tạo JWT, cấp refresh token qua cookie HttpOnly và trả HTTP 201; giao diện thiết lập phiên đăng nhập ngay. CAPTCHA không hợp lệ hoặc dịch vụ xác minh không khả dụng sẽ chặn đăng ký.',
 [m(0,1,'Nhập thông tin và xác nhận mật khẩu'),m(1,2,'Gửi yêu cầu đăng ký'),ref('Kiểm tra dữ liệu và CAPTCHA; dừng nếu không đạt'),m(2,3,'Kiểm tra email, lấy role Customer'),r(3,2,'Kết quả kiểm tra'),alt('Email đã tồn tại'),r(2,1,'Báo email đã được sử dụng'),other('Email chưa được sử dụng'),m(2,2,'Băm mật khẩu, tạo User và UserRole'),m(2,3,'Lưu tài khoản và vai trò'),r(3,2,'Đã lưu'),ref('TokenService tạo JWT; cấp refresh cookie'),r(2,1,'201 và thông tin phiên'),end(),r(1,0,'Thông báo và điều hướng')],service='AuthController\n/ AuthService')
add('Đăng nhập tài khoản','R1.2, R1.3','Người dùng đã có tài khoản.',
 'AuthService tìm người dùng kèm vai trò, kiểm tra khóa tạm, xác minh mật khẩu BCrypt và trạng thái hoạt động. Khi sai mật khẩu, bộ đếm thất bại được cập nhật; đủ ngưỡng thì khóa tạm. Tài khoản bật xác thực hai lớp phải hoàn tất bước xác minh trước khi được cấp token. TokenService chỉ tạo JWT và thời hạn; AuthService ghép thông tin người dùng vào AuthResponseDto.',
 [m(0,1,'Nhập email, mật khẩu'),m(1,2,'Gửi yêu cầu đăng nhập'),ref('Kiểm tra dữ liệu và CAPTCHA'),m(2,3,'Tìm User và các vai trò'),r(3,2,'Thông tin tài khoản'),m(2,2,'Kiểm tra khóa, BCrypt, IsActive'),alt('Sai thông tin hoặc tài khoản bị khóa'),m(2,3,'Ghi số lần sai nếu áp dụng'),r(2,1,'401 và thông báo lỗi'),other('Hợp lệ và bật xác thực hai lớp'),ref('Thực hiện Hình 2.2.1-04 trước khi cấp phiên'),other('Hợp lệ và không bật xác thực hai lớp'),ref('Reset số lần sai; tạo JWT và refresh cookie'),r(2,1,'200 và AuthResponseDto'),end(),r(1,0,'Hiển thị kết quả, điều hướng')],service='AuthController\n/ AuthService')
add('Xác minh đăng nhập hai lớp','Mở rộng R1.2','Mật khẩu đúng và tài khoản đã bật xác thực hai lớp.',
 'Hệ thống trả challenge thay vì JWT. Người dùng nhập mã TOTP hoặc mã khôi phục dùng một lần. TwoFactorService kiểm tra challenge, trạng thái khóa và mã xác minh. Chỉ khi hợp lệ, hệ thống mới tạo phiên đăng nhập đầy đủ. Challenge hết hạn hoặc tài khoản bị khóa yêu cầu quay lại bước đăng nhập.',
 [m(2,3,'Tạo challenge đăng nhập'),r(3,2,'Challenge đã lưu'),r(2,1,'requiresTwoFactor và challenge'),r(1,0,'Yêu cầu mã xác minh'),m(0,1,'Nhập TOTP hoặc mã khôi phục'),m(1,2,'Gửi challenge và mã'),m(2,3,'Đọc challenge, tài khoản, mã khôi phục'),r(3,2,'Dữ liệu xác minh'),alt('Mã / challenge không hợp lệ'),m(2,3,'Ghi nhận thất bại nếu áp dụng'),r(2,1,'401 và lý do'),other('Hợp lệ'),m(2,3,'Tiêu thụ challenge / mã, reset số lần sai'),ref('Tạo JWT, cấp refresh cookie'),r(2,1,'200 và phiên đăng nhập'),end()],service='AuthController\n/ TwoFactorService')
add('Quên mật khẩu','R1.4','Người dùng mở chức năng quên mật khẩu.',
 'Người dùng gửi email. Với tài khoản đang hoạt động, hệ thống tạo token đặt lại mật khẩu có hạn 30 phút, lưu token và gửi liên kết qua EmailService. Phản hồi không tiết lộ email có tồn tại hay không. Việc đặt lại mật khẩu được trình bày ở sơ đồ kế tiếp.',
 [m(0,1,'Nhập email'),m(1,2,'Yêu cầu quên mật khẩu'),m(2,3,'Tìm tài khoản theo email'),r(3,2,'Tài khoản hoặc không có'),opt('Tài khoản tồn tại và hoạt động'),m(2,2,'Tạo token, hạn dùng 30 phút'),m(2,3,'Lưu token đặt lại mật khẩu'),m(2,4,'Gửi email chứa liên kết'),r(4,2,'Kết quả gửi'),end(),r(2,1,'Thông báo đã tiếp nhận yêu cầu'),r(1,0,'Hướng dẫn kiểm tra email')],lanes=['Người dùng','Angular','AuthService','EF Core / SQL','EmailService'])
add('Đặt lại mật khẩu','R1.4','Người dùng mở liên kết đặt lại mật khẩu và nhập mật khẩu mới.',
 'AuthService kiểm tra token và thời hạn. Nếu hợp lệ, hệ thống băm mật khẩu mới, xóa token đặt lại và thu hồi toàn bộ refresh token trong cùng giao dịch. Token đặt lại chỉ dùng được một lần. Người dùng đăng nhập lại bằng mật khẩu mới.',
 [m(0,1,'Nhập và xác nhận mật khẩu mới'),m(1,2,'Gửi token và mật khẩu mới'),m(2,3,'Tìm tài khoản theo token'),r(3,2,'Token và thời hạn'),alt('Token sai hoặc hết hạn'),r(2,1,'Từ chối đặt lại mật khẩu'),other('Token còn hiệu lực'),m(2,2,'Băm mật khẩu mới bằng BCrypt'),m(2,3,'Bắt đầu giao dịch; lưu mật khẩu, xóa token'),ref('RefreshTokenService thu hồi các phiên'),m(2,3,'Commit giao dịch'),r(2,1,'Đặt lại mật khẩu thành công'),end(),r(1,0,'Thông báo kết quả')],service='AuthController\n/ AuthService')
add('Làm mới phiên và đăng xuất','R1.6 và duy trì phiên','Trình duyệt có hoặc đã có refresh cookie.',
 'Access token được giữ trong bộ nhớ giao diện; refresh token nằm trong cookie HttpOnly. Khi làm mới phiên, máy chủ kiểm tra và xoay vòng refresh token rồi cấp JWT mới. Khi đăng xuất, máy chủ thu hồi họ refresh token của phiên và xóa cookie; giao diện xóa trạng thái đăng nhập. Hai thao tác là hai lựa chọn độc lập.',
 [alt('Làm mới phiên'),m(1,2,'POST refresh kèm cookie'),m(2,3,'Kiểm tra và xoay vòng refresh token'),r(3,2,'Kết quả xác thực phiên'),ref('Nếu hợp lệ: tạo JWT và đặt cookie mới; nếu không: trả 401'),r(2,1,'Phiên mới hoặc yêu cầu đăng nhập lại'),other('Đăng xuất'),m(0,1,'Chọn đăng xuất'),m(1,2,'POST logout kèm cookie'),m(2,3,'Thu hồi họ refresh token'),r(2,1,'Xóa cookie, trả kết quả'),m(1,1,'Xóa token trong bộ nhớ và thông tin phiên'),r(1,0,'Chuyển về trạng thái chưa đăng nhập'),end()],service='AuthController\n/ RefreshTokenService')
add('Xem và tìm kiếm sản phẩm','R5','Khách truy cập trang sản phẩm; không bắt buộc đăng nhập.',
 'Giao diện gửi từ khóa, bộ lọc, thứ tự sắp xếp và thông tin phân trang. StorefrontService truy vấn các sản phẩm đang kinh doanh. Khi xem chi tiết, hệ thống tải thông tin sản phẩm, biến thể, hình ảnh và đánh giá đã duyệt. Không có kết quả tìm kiếm là danh sách rỗng; sản phẩm không tồn tại hoặc bị ẩn thì trả lỗi phù hợp.',
 [alt('Xem danh sách / tìm kiếm / lọc'),m(0,1,'Nhập từ khóa, chọn bộ lọc'),m(1,2,'Gửi điều kiện và trang cần xem'),m(2,3,'Truy vấn sản phẩm đang kinh doanh'),r(3,2,'Danh sách và tổng số kết quả'),r(2,1,'Kết quả đã phân trang'),r(1,0,'Hiển thị sản phẩm hoặc danh sách rỗng'),other('Xem chi tiết'),m(0,1,'Chọn một sản phẩm'),m(1,2,'Yêu cầu chi tiết theo slug'),m(2,3,'Đọc sản phẩm, biến thể, ảnh, đánh giá'),r(3,2,'Dữ liệu sản phẩm'),r(2,1,'Chi tiết hoặc lỗi không tìm thấy'),r(1,0,'Hiển thị thông tin'),end()],service='ProductsController\n/ StorefrontService')
add('Quản lý giỏ hàng','R6','Khách hàng đã đăng nhập.',
 'Mỗi dòng giỏ hàng gắn với một biến thể theo màu và kích thước. Khi thêm hoặc cập nhật, CartService kiểm tra trạng thái kinh doanh và số lượng; biến thể đã có thì cộng dồn. Xóa một dòng hoặc toàn bộ giỏ cần xác nhận trên giao diện. Xem giỏ trả số lượng, giá, tổng tiền và cảnh báo tồn kho.',
 [m(0,1,'Chọn thao tác trên giỏ hàng'),m(1,2,'Gửi thao tác và dữ liệu'),m(2,3,'Đọc giỏ thuộc user và biến thể'),r(3,2,'Dữ liệu giỏ, giá, tồn kho'),alt('Thêm / cập nhật hợp lệ'),m(2,2,'Kiểm tra số lượng; cộng dồn nếu thêm'),m(2,3,'Lưu dòng giỏ hàng'),other('Xóa sau khi xác nhận'),m(2,3,'Xóa dòng hoặc toàn bộ giỏ của user'),other('Chỉ xem giỏ'),m(2,2,'Tính thành tiền, tạm tính và cảnh báo'),other('Dữ liệu hoặc tồn kho không hợp lệ'),r(2,1,'Trả lỗi, không cập nhật'),end(),r(2,1,'Kết quả thao tác nếu thành công'),r(1,0,'Cập nhật giỏ hoặc hiển thị lỗi')],service='CartController\n/ CartService')
add('Quản lý danh sách yêu thích','R7','Khách hàng đã đăng nhập.',
 'WishlistService xác định người dùng từ phiên đăng nhập. Khách hàng có thể xem, thêm hoặc bỏ sản phẩm yêu thích. Hệ thống kiểm tra sản phẩm khi thêm và chỉ thao tác trên danh sách thuộc người dùng hiện tại.',
 [m(0,1,'Xem / thêm / bỏ yêu thích'),m(1,2,'Gửi thao tác yêu thích'),m(2,3,'Đọc sản phẩm và mục yêu thích của user'),r(3,2,'Kết quả kiểm tra'),alt('Thêm sản phẩm hợp lệ'),m(2,3,'Lưu mục yêu thích nếu chưa có'),other('Bỏ yêu thích'),m(2,3,'Xóa mục yêu thích của user'),other('Xem danh sách'),m(2,3,'Truy vấn danh sách và thông tin sản phẩm'),end(),r(2,1,'Kết quả hoặc lỗi sản phẩm'),r(1,0,'Hiển thị trạng thái yêu thích')],service='WishlistController\n/ WishlistService')
add('Cập nhật hồ sơ và đổi mật khẩu','R8.1, R1.5','Người dùng đã đăng nhập.',
 'Người dùng được cập nhật họ tên, số điện thoại, ảnh đại diện; email không thay đổi qua chức năng này. Đổi mật khẩu yêu cầu xác minh mật khẩu hiện tại trước khi băm mật khẩu mới. Các phiên refresh được thu hồi theo luồng đổi mật khẩu của máy chủ; giao diện yêu cầu đăng nhập lại.',
 [m(0,1,'Nhập thông tin muốn thay đổi'),m(1,2,'Gửi yêu cầu tài khoản hiện tại'),m(2,3,'Đọc người dùng từ userId đã xác thực'),r(3,2,'Thông tin tài khoản'),alt('Cập nhật hồ sơ hợp lệ'),m(2,3,'Lưu họ tên, điện thoại, đường dẫn ảnh'),r(2,1,'Hồ sơ đã cập nhật'),other('Đổi mật khẩu'),m(2,2,'Xác minh mật khẩu hiện tại bằng BCrypt'),alt('Mật khẩu hiện tại đúng'),m(2,3,'Lưu mật khẩu mới; thu hồi refresh token'),r(2,1,'Đổi mật khẩu thành công'),other('Mật khẩu hiện tại sai'),r(2,1,'Thông báo lỗi, không thay đổi'),end(),end(),r(1,0,'Hiển thị kết quả')],service='UsersController\n/ UserProfileService')
add('Quản lý địa chỉ giao hàng','R8.2, R8.3','Khách hàng đã đăng nhập.',
 'Địa chỉ phải thuộc tài khoản hiện tại. Khi thêm, sửa hoặc chọn mặc định, hệ thống kiểm tra người nhận, số điện thoại và các trường địa chỉ. Nếu đặt một địa chỉ làm mặc định, các địa chỉ khác của cùng người dùng được bỏ trạng thái mặc định. Dữ liệu địa chỉ được dùng khi tạo đơn hàng.',
 [m(0,1,'Xem / thêm / sửa / xóa / chọn mặc định'),m(1,2,'Gửi thao tác địa chỉ'),m(2,3,'Đọc địa chỉ của người dùng'),r(3,2,'Danh sách hoặc địa chỉ được chọn'),alt('Dữ liệu hoặc quyền sở hữu không hợp lệ'),r(2,1,'Thông báo lỗi'),other('Thao tác hợp lệ'),opt('Chọn địa chỉ mặc định'),m(2,3,'Bỏ mặc định các địa chỉ khác của user'),end(),m(2,3,'Lưu thay đổi hoặc truy vấn danh sách'),r(3,2,'Kết quả'),r(2,1,'Địa chỉ sau thao tác'),end(),r(1,0,'Hiển thị sổ địa chỉ')],service='AddressesController\n/ AddressService')
add('Kiểm tra và áp dụng mã giảm giá','R9.4, R9.5','Khách hàng đã đăng nhập, có sản phẩm cần mua.',
 'CouponService đọc mã, lịch sử sử dụng và các mặt hàng cần mua để kiểm tra thời gian, trạng thái, lượt dùng, phạm vi và giá trị tối thiểu. Hệ thống trả số tiền giảm hoặc lý do không áp dụng được. Kiểm tra mã chỉ cho kết quả dự kiến; lượt sử dụng chỉ được ghi khi tạo đơn thành công và mã phải được kiểm tra lại tại thời điểm đó.',
 [m(0,1,'Xem mã hoặc nhập mã giảm giá'),m(1,2,'Yêu cầu kiểm tra mã với giỏ / mua ngay'),m(2,3,'Đọc coupon, phạm vi, lượt dùng, hàng mua'),r(3,2,'Dữ liệu kiểm tra'),m(2,2,'Kiểm tra điều kiện và tính tiền giảm'),alt('Mã áp dụng được'),r(2,1,'Số tiền giảm và kết quả kiểm tra'),m(1,1,'Cập nhật tổng tiền dự kiến'),other('Mã không áp dụng được'),r(2,1,'Lý do từ chối'),end(),r(1,0,'Hiển thị kết quả áp dụng')],service='CouponsController\n/ CouponService')
add('Tạo đơn hàng','R10.1–R10.3','Khách hàng đăng nhập, có địa chỉ hợp lệ và hàng cần mua.',
 'OrderService mở giao dịch, kiểm tra địa chỉ thuộc khách hàng và lấy giỏ hàng hoặc mặt hàng mua ngay. Máy chủ kiểm tra tồn kho, giá và mã giảm giá, sau đó tự tính tổng tiền. Đơn, chi tiết, thanh toán ban đầu, lịch sử, tồn kho và lượt dùng mã được lưu cùng giao dịch. Mua từ giỏ sẽ xóa các dòng giỏ; mua ngay giữ nguyên giỏ. Lỗi trước commit làm giao dịch bị hủy.',
 [m(0,1,'Xác nhận địa chỉ, hàng và phương thức trả tiền'),m(1,2,'Gửi yêu cầu đặt hàng'),m(2,3,'Begin transaction; đọc địa chỉ và hàng mua'),r(3,2,'Dữ liệu đặt hàng'),ref('Kiểm tra sản phẩm, tồn kho; kiểm tra lại coupon nếu có'),alt('Không đủ điều kiện tạo đơn'),m(2,3,'Rollback khi kết thúc giao dịch lỗi'),r(2,1,'Báo lý do, không tạo đơn'),other('Hợp lệ'),m(2,2,'Tính tiền; tạo bản chụp thông tin đơn'),m(2,3,'Lưu đơn, thanh toán, lịch sử; trừ kho; ghi coupon'),m(2,3,'Xóa giỏ nếu mua từ giỏ; SaveChanges; commit'),r(3,2,'Đơn đã lưu thành công'),r(2,1,'Chi tiết đơn hàng mới'),end(),r(1,0,'Hiển thị kết quả đặt hàng')],service='OrdersController\n/ OrderService')
add('Thanh toán QR mô phỏng','R10.4','Đơn thuộc khách hàng và sử dụng phương thức MockQR.',
 'Hệ thống trả mã đơn, số tiền và nội dung chuyển khoản để giao diện hiển thị QR mô phỏng. Khi khách hàng xác nhận, máy chủ kiểm tra quyền sở hữu, phương thức thanh toán và trạng thái đơn rồi cập nhật Payment cùng Order sang Paid. Đây là xác nhận trong mô phỏng, không phải callback xác thực của ngân hàng. COD được ghi nhận khi quản trị viên hoàn thành giao hàng.',
 [m(0,1,'Mở thanh toán QR của đơn'),m(1,2,'Yêu cầu thông tin QR'),m(2,3,'Đọc đơn và phương thức thanh toán'),r(3,2,'Dữ liệu đơn'),r(2,1,'Mã đơn, số tiền, nội dung chuyển khoản'),r(1,0,'Hiển thị QR mô phỏng'),m(0,1,'Xác nhận đã thanh toán'),m(1,2,'Gửi xác nhận MockQR'),alt('Đơn hợp lệ và chưa bị hủy'),m(2,3,'Cập nhật Payment và Order thành Paid'),r(2,1,'Kết quả thanh toán'),other('Không hợp lệ'),r(2,1,'Từ chối xác nhận'),end(),r(1,0,'Hiển thị trạng thái thanh toán')],service='OrdersController\n/ OrderService')
add('Theo dõi và hủy đơn hàng','R11.1, R11.2','Khách hàng đã đăng nhập và chỉ thao tác trên đơn của mình.',
 'Khi xem đơn, hệ thống trả danh sách hoặc chi tiết kèm lịch sử trạng thái. Khi hủy, OrderService chỉ cho phép đơn Pending và chưa Paid. Hệ thống hoàn tồn kho, hoàn lượt dùng mã nếu có, chuyển trạng thái Cancelled và ghi lịch sử trong cùng giao dịch. Không tìm thấy hoặc không sở hữu đơn thì không cho truy cập.',
 [m(0,1,'Xem đơn hoặc yêu cầu hủy'),m(1,2,'Gửi thao tác đơn hàng'),m(2,3,'Đọc đơn thuộc user và dữ liệu liên quan'),r(3,2,'Đơn và lịch sử'),alt('Xem đơn'),r(2,1,'Danh sách / chi tiết đơn'),other('Hủy: Pending và chưa Paid'),m(2,3,'Trong giao dịch: hoàn kho, hoàn lượt coupon'),m(2,3,'Lưu Cancelled, lịch sử và commit'),r(2,1,'Hủy thành công'),other('Không đủ điều kiện hủy / không tìm thấy'),r(2,1,'Từ chối và nêu lý do'),end(),r(1,0,'Hiển thị kết quả')],service='OrdersController\n/ OrderService')
add('Gửi và quản lý đánh giá của khách hàng','R12.1, R12.2','Khách hàng đã đăng nhập; đánh giá mới gắn với dòng đơn đã hoàn thành.',
 'ReviewService kiểm tra dòng đơn thuộc người dùng, đơn Completed và chưa có đánh giá cho OrderItemId đó. Đánh giá mới có IsApproved = false; đánh giá sửa cũng phải duyệt lại. Khách hàng được xem, sửa hoặc xóa đánh giá của mình; điểm đánh giá từ 1 đến 5.',
 [m(0,1,'Gửi / sửa / xóa / xem đánh giá'),m(1,2,'Gửi nội dung và định danh liên quan'),m(2,3,'Đọc dòng đơn hoặc đánh giá của user'),r(3,2,'Dữ liệu kiểm tra'),m(2,2,'Kiểm tra quyền, điều kiện và điểm số'),alt('Không đủ điều kiện'),r(2,1,'Thông báo lỗi'),other('Tạo hoặc sửa hợp lệ'),m(2,3,'Lưu đánh giá với IsApproved = false'),r(2,1,'Đánh giá đang chờ duyệt'),other('Xóa hoặc xem hợp lệ'),m(2,3,'Xóa / truy vấn đánh giá của user'),r(2,1,'Kết quả thao tác'),end(),r(1,0,'Hiển thị đánh giá của tôi')],service='ReviewsController\n/ ReviewService')
add('Quản trị sản phẩm và biến thể','R2','Quản trị viên đã đăng nhập và có quyền Admin.',
 'ProductService kiểm tra danh mục, thương hiệu, dữ liệu giá và slug. Với biến thể, hệ thống kiểm tra màu, kích thước, SKU và tổ hợp biến thể không trùng. Tồn kho được quản lý qua nhập kho và bán hàng. Thao tác xóa sản phẩm hoặc biến thể được thực hiện bằng cách ngừng hoạt động để giữ liên kết dữ liệu.',
 [m(0,1,'Nhập sản phẩm / biến thể hoặc chọn ngừng bán'),m(1,2,'Gửi yêu cầu quản trị'),m(2,3,'Đọc sản phẩm, dữ liệu tham chiếu, SKU'),r(3,2,'Dữ liệu kiểm tra'),m(2,2,'Kiểm tra giá, slug, SKU và tổ hợp biến thể'),alt('Không hợp lệ hoặc trùng dữ liệu'),r(2,1,'Thông báo lỗi'),other('Thêm / sửa hợp lệ'),m(2,3,'Lưu Product hoặc ProductVariant'),r(2,1,'Dữ liệu đã cập nhật'),other('Ngừng bán'),m(2,3,'Đặt IsActive = false'),r(2,1,'Xác nhận ngừng bán'),end(),r(1,0,'Hiển thị kết quả')],service='AdminProductsController\n/ ProductService',actor='Quản trị viên')
add('Quản lý hình ảnh sản phẩm','R2.3','Quản trị viên có quyền Admin và sản phẩm đã tồn tại.',
 'Tệp ảnh được kiểm tra định dạng và kích thước trước khi lưu qua chức năng upload. Sau đó ProductService lưu đường dẫn ảnh, thứ tự và trạng thái ảnh chính. Sửa, xóa ảnh hoặc đổi ảnh đại diện cập nhật các bản ghi ProductImage tương ứng. Lưu tệp và lưu liên kết ảnh là hai thao tác riêng.',
 [m(0,1,'Chọn ảnh sản phẩm'),m(1,2,'Gửi tệp JPG / PNG / WEBP'),m(2,2,'Kiểm tra loại tệp và giới hạn 5 MB'),alt('Tệp không hợp lệ'),r(2,1,'Từ chối tải ảnh'),other('Tệp hợp lệ'),m(2,3,'Lưu tệp ảnh'),r(3,2,'Đường dẫn ảnh'),r(2,1,'Kết quả upload'),m(1,2,'Gửi đường dẫn, thứ tự, IsMain'),m(2,3,'Lưu ProductImage, điều chỉnh ảnh chính'),r(2,1,'Danh sách ảnh đã cập nhật'),end(),r(1,0,'Hiển thị kết quả')],service='Uploads / AdminProducts\nController + Service',actor='Quản trị viên',lanes=['Quản trị viên','Angular','API quản lý ảnh','Tệp ảnh /\nEF Core / SQL'])
add('Quản lý dữ liệu danh mục và nhà cung cấp','R3, R4.1','Quản trị viên có quyền Admin.',
 'Các chức năng danh mục, thương hiệu, màu, kích thước và nhà cung cấp có luồng chung: nhận dữ liệu, kiểm tra ràng buộc, truy vấn hoặc lưu và trả kết quả. Service tương ứng kiểm tra các liên kết trước thao tác xóa; tùy loại dữ liệu, hệ thống từ chối xóa hoặc chuyển sang trạng thái ngừng hoạt động. Không xóa tùy ý dữ liệu đã được sử dụng.',
 [m(0,1,'Chọn loại dữ liệu và thao tác'),m(1,2,'Gửi yêu cầu đến API tương ứng'),m(2,3,'Đọc bản ghi và dữ liệu liên quan'),r(3,2,'Kết quả truy vấn'),m(2,2,'Kiểm tra tên, mã và ràng buộc liên kết'),alt('Vi phạm ràng buộc'),r(2,1,'Báo lỗi và giữ nguyên dữ liệu'),other('Hợp lệ'),m(2,3,'Thêm / sửa / xóa / đổi trạng thái theo nghiệp vụ'),r(3,2,'Kết quả lưu'),r(2,1,'Danh sách hoặc bản ghi đã cập nhật'),end(),r(1,0,'Hiển thị kết quả')],service='Controller / Service\ncủa nhóm danh mục',actor='Quản trị viên')
add('Lập và điều chỉnh phiếu nhập kho','R4.2–R4.4','Quản trị viên có quyền Admin; nhà cung cấp và biến thể đã tồn tại.',
 'Khi lập phiếu, hệ thống cộng tồn và tính giá vốn bình quân gia quyền. Sửa phiếu hoàn tác phần nhập cũ rồi áp dụng nội dung mới; hủy phiếu cần lý do và hoàn tác lượng nhập. Hệ thống từ chối nếu tồn sau điều chỉnh âm hoặc phiếu đã hủy. Các thay đổi của mỗi thao tác được lưu bằng một lần SaveChangesAsync.',
 [m(0,1,'Lập / sửa / hủy phiếu nhập'),m(1,2,'Gửi nhà cung cấp và các dòng hàng'),m(2,3,'Đọc phiếu, nhà cung cấp, biến thể, tồn kho'),r(3,2,'Dữ liệu hiện tại'),m(2,2,'Kiểm tra điều kiện và tồn sau điều chỉnh'),alt('Phiếu đã hủy hoặc tồn không đủ'),r(2,1,'Từ chối điều chỉnh'),other('Hợp lệ'),ref('Tạo: cộng kho; sửa: hoàn tác rồi nhập lại; hủy: hoàn tác'),m(2,2,'Tính lại giá vốn và tổng phiếu'),m(2,3,'Lưu phiếu, dòng hàng và tồn kho'),r(2,1,'Chi tiết phiếu sau xử lý'),end(),r(1,0,'Hiển thị phiếu nhập')],service='AdminPurchaseOrders\n/ PurchaseOrderService',actor='Quản trị viên')
add('Quản trị mã giảm giá','R9.1–R9.3','Quản trị viên có quyền Admin.',
 'Quản trị viên khai báo điều kiện áp dụng, thời gian, giới hạn lượt và phạm vi sản phẩm hoặc danh mục. CouponService kiểm tra dữ liệu, lưu coupon cùng các liên kết phạm vi. Mã giảm giá được bật hoặc tắt thay vì xóa để giữ lịch sử. Xem lịch sử truy vấn CouponUsage và các thông tin liên quan.',
 [m(0,1,'Tạo / sửa / bật tắt / xem lịch sử mã'),m(1,2,'Gửi yêu cầu quản trị mã'),m(2,3,'Đọc coupon và dữ liệu phạm vi'),r(3,2,'Dữ liệu kiểm tra'),alt('Tạo hoặc sửa'),m(2,2,'Kiểm tra mã, giá trị, thời gian, phạm vi'),ref('Nếu không hợp lệ: trả lỗi và kết thúc nhánh'),m(2,3,'Lưu coupon và liên kết phạm vi'),other('Bật hoặc tắt'),m(2,3,'Cập nhật IsActive'),other('Xem lịch sử sử dụng'),m(2,3,'Truy vấn CouponUsage và đơn liên quan'),end(),r(2,1,'Kết quả thao tác'),r(1,0,'Hiển thị thông tin mã')],service='AdminCouponsController\n/ CouponService',actor='Quản trị viên')
add('Quản trị trạng thái đơn hàng','R11.3, R11.4, R11.6','Quản trị viên có quyền Admin và chọn một đơn hàng.',
 'Luồng xử lý chính là Pending → Confirmed → Shipping → Completed. Hủy được ở Pending hoặc Confirmed; khi hủy phải hoàn kho và lượt coupon, đồng thời ghi nhận hoàn tiền nếu đã Paid. Khi Completed, COD tự được ghi Paid; đơn QR phải thanh toán trước. Cập nhật trạng thái và lịch sử được thực hiện trong giao dịch.',
 [m(0,1,'Chọn trạng thái mới và ghi chú'),m(1,2,'Yêu cầu cập nhật đơn'),m(2,3,'Begin transaction; đọc đơn và dữ liệu liên quan'),r(3,2,'Đơn hiện tại'),m(2,2,'Kiểm tra chuyển trạng thái và thanh toán'),alt('Không được phép chuyển'),r(2,1,'Từ chối; rollback giao dịch'),other('Chuyển hợp lệ'),ref('Hủy: hoàn kho / coupon / ghi hoàn tiền; hoàn thành COD: ghi Paid'),m(2,3,'Lưu trạng thái và lịch sử người thực hiện'),m(2,3,'SaveChanges và commit'),r(2,1,'Chi tiết đơn đã cập nhật'),end(),r(1,0,'Hiển thị tiến trình xử lý')],service='AdminOrdersController\n/ AdminOrderService',actor='Quản trị viên')
add('Quản trị trạng thái thanh toán','R11.5','Quản trị viên có quyền Admin và chọn đơn có bản ghi thanh toán.',
 'Máy chủ kiểm tra trạng thái đơn và phương thức thanh toán trước khi thay đổi Payment. Đơn đã hủy hoặc đã trả không được cập nhật qua luồng này; không đặt Refunded trực tiếp. COD chỉ được ghi Paid khi đơn Completed. Thông tin Payment và PaymentStatus trên Order được cập nhật đồng bộ.',
 [m(0,1,'Chọn trạng thái thanh toán'),m(1,2,'Gửi yêu cầu cập nhật thanh toán'),m(2,3,'Đọc đơn và Payment'),r(3,2,'Trạng thái và phương thức thanh toán'),m(2,2,'Kiểm tra điều kiện cập nhật'),alt('Vi phạm điều kiện'),r(2,1,'Từ chối và nêu lý do'),other('Hợp lệ'),m(2,3,'Lưu trạng thái Payment và Order'),r(3,2,'Kết quả cập nhật'),r(2,1,'Thông tin thanh toán mới'),end(),r(1,0,'Hiển thị kết quả')],service='AdminOrdersController\n/ AdminOrderService',actor='Quản trị viên')
add('Kiểm duyệt đánh giá sản phẩm','R12.3','Quản trị viên có quyền Admin.',
 'Quản trị viên xem danh sách đánh giá và chọn duyệt hoặc ẩn. ReviewService kiểm tra đánh giá tồn tại rồi cập nhật IsApproved. Trang sản phẩm chỉ truy vấn các đánh giá đã duyệt, nên đánh giá bị ẩn không xuất hiện công khai.',
 [m(0,1,'Mở danh sách đánh giá'),m(1,2,'Yêu cầu danh sách quản trị'),m(2,3,'Truy vấn đánh giá theo bộ lọc'),r(3,2,'Danh sách đánh giá'),r(2,1,'Kết quả phân trang'),m(0,1,'Chọn duyệt hoặc ẩn'),m(1,2,'Gửi reviewId và IsApproved'),m(2,3,'Kiểm tra và cập nhật đánh giá'),alt('Đánh giá tồn tại'),r(2,1,'Trạng thái duyệt đã cập nhật'),other('Không tồn tại'),r(2,1,'Lỗi không tìm thấy'),end(),r(1,0,'Hiển thị kết quả')],service='AdminReviewsController\n/ ReviewService',actor='Quản trị viên')
add('Quản lý tài khoản người dùng','R13','Quản trị viên có quyền Admin.',
 'Quản trị viên tìm kiếm, lọc và xem chi tiết tài khoản. Khi thay đổi trạng thái, AdminUserService kiểm tra tài khoản tồn tại và ngăn quản trị viên tự khóa chính mình. Tài khoản bị vô hiệu hóa không thể đăng nhập; các refresh token của tài khoản bị khóa được thu hồi theo xử lý của Service.',
 [m(0,1,'Tìm kiếm / xem tài khoản'),m(1,2,'Gửi bộ lọc hoặc userId'),m(2,3,'Truy vấn người dùng và vai trò'),r(3,2,'Dữ liệu tài khoản'),r(2,1,'Danh sách / chi tiết'),m(0,1,'Chọn khóa hoặc mở khóa'),m(1,2,'Gửi trạng thái hoạt động mới'),m(2,2,'Kiểm tra tài khoản và không tự khóa'),alt('Không hợp lệ'),r(2,1,'Từ chối thay đổi'),other('Hợp lệ'),m(2,3,'Lưu IsActive; thu hồi phiên nếu khóa'),r(2,1,'Trạng thái mới'),end(),r(1,0,'Hiển thị kết quả')],service='AdminUsersController\n/ AdminUserService',actor='Quản trị viên')
add('Xem thống kê quản trị','R14','Quản trị viên có quyền Admin.',
 'AdminDashboardService kiểm tra khoảng ngày, chuyển mốc ngày Việt Nam sang UTC để lọc thời điểm tạo đơn. Doanh thu và giá vốn được tính từ đơn Completed; lợi nhuận bằng doanh thu trừ giá vốn. Dashboard còn trả tổng đơn, đơn chờ, khách hàng, hàng sắp hết và sản phẩm bán chạy. Tổng khách hàng và danh sách tồn thấp phản ánh dữ liệu hiện tại, không cùng phạm vi ngày với mọi chỉ số đơn hàng.',
 [m(0,1,'Mở dashboard, chọn khoảng ngày'),m(1,2,'Gửi yêu cầu thống kê'),alt('Ngày bắt đầu lớn hơn ngày kết thúc'),r(2,1,'Báo khoảng thời gian không hợp lệ'),other('Khoảng ngày hợp lệ'),m(2,2,'Chuẩn hóa mốc ngày Việt Nam sang UTC'),m(2,3,'Truy vấn đơn, chi tiết, khách hàng, tồn kho'),r(3,2,'Các tập dữ liệu thống kê'),m(2,2,'Tính doanh thu, giá vốn, lợi nhuận và tỷ suất'),r(2,1,'DashboardDto và các danh sách'),end(),r(1,0,'Hiển thị chỉ số và biểu đồ')],service='AdminDashboardController\n/ AdminDashboardService',actor='Quản trị viên')

def wrap(text,width,font=F):
 lines=[]
 for line in text.split('\n'):
  cur=''
  for word in line.split():
   if font.getlength((cur+' '+word).strip())>width and cur:
    lines.append(cur);cur=word
   else:cur=(cur+' '+word).strip()
  lines.append(cur)
 return lines

def render(s,number):
 xs=[180+(W-360)*i/(len(s['lanes'])-1) for i in range(len(s['lanes']))]
 layout=[];y=160;stack=[];frames=[];num=0
 for e in s['events']:
  if e[0] in ('m','r'):
   num+=1;a,b=e[1:3];width=abs(xs[b]-xs[a])-34 if a!=b else 405
   lines=wrap(f'{num}. '+e[3],width)
   h=max(88,len(lines)*32+38)+(26 if a==b else 0)
   layout.append((e,y,h,lines));y+=h
  elif e[0]=='start':
   stack.append((y,e[1],len(stack)));lines=wrap('['+e[2]+']',W-190,FS);h=len(lines)*31+24
   layout.append((e,y,h,lines));y+=h
  elif e[0]=='else':
   lines=wrap('['+e[1]+']',W-190,FS);h=len(lines)*31+22
   layout.append((e,y,h,lines));y+=h
  elif e[0]=='end':
   top,kind,dep=stack.pop();frames.append((top,y+8,kind,dep));y+=28
  else:
   lines=wrap(e[1],W-180,FS);h=len(lines)*31+30
   layout.append((e,y,h,lines));y+=h+13
 assert not stack
 im=Image.new('RGB',(W,int(y+35)),'white');d=ImageDraw.Draw(im)
 for x in xs:
  for yy in range(120,int(y),15):d.line((x,yy,x,yy+7),fill='#999999',width=2)
 for i,(x,label) in enumerate(zip(xs,s['lanes'])):
  header_font=FB
  for font_size in range(28,17,-1):
   header_font=ImageFont.truetype(BOLD,font_size)
   lines=wrap(label,290,header_font)
   if len(lines)<=2 and max(header_font.getlength(t) for t in lines)<=290:break
  d.rectangle((x-157,12,x+157,112),fill='white',outline='black',width=2)
  for j,t in enumerate(lines):d.text((x,35+j*32),t,font=header_font,fill='black',anchor='mt')
 for top,bottom,kind,dep in frames:
  left=18+dep*14;right=W-18-dep*14
  d.rectangle((left,top,right,bottom),outline='#333333',width=2)
  d.rectangle((left,top,left+80,top+38),fill='white',outline='black',width=2)
  d.text((left+10,top+3),kind,font=FB,fill='black')
 for e,yy,h,lines in layout:
  if e[0] in ('m','r'):
   a,b=e[1:3];xa,xb=xs[a],xs[b];liney=yy+h-18
   if a==b:
    direction=1 if a<len(xs)-1 else -1
    xend=xa+110*direction
    d.line([(xa,liney-24),(xend,liney-24),(xend,liney),(xa,liney)],fill='black',width=2)
    tx=xa+direction*220
   else:
    tx=(xa+xb)/2
    if e[0]=='r':
     for xx in range(int(min(xa,xb)),int(max(xa,xb)),16):d.line((xx,liney,min(xx+8,max(xa,xb)),liney),fill='black',width=2)
    else:d.line((xa,liney,xb,liney),fill='black',width=2)
   direction=1 if xb>xa else -1
   if a==b:direction= -1 if a<len(xs)-1 else 1
   if e[0]=='m':d.polygon([(xb,liney),(xb-direction*14,liney-7),(xb-direction*14,liney+7)],fill='black')
   else:d.line([(xb-direction*14,liney-7),(xb,liney),(xb-direction*14,liney+7)],fill='black',width=2)
   for j,t in enumerate(lines):
    bb=d.textbbox((tx,yy+4+j*32),t,font=F,anchor='mt');d.rectangle((bb[0]-3,bb[1]-1,bb[2]+3,bb[3]+1),fill='white');d.text((tx,yy+4+j*32),t,font=F,fill='black',anchor='mt')
  elif e[0] in ('start','else'):
   if e[0]=='else':
    for xx in range(20,W-20,16):d.line((xx,yy,xx+8,yy),fill='#555555',width=2)
   for j,t in enumerate(lines):
    x=115;bb=d.textbbox((x,yy+7+j*31),t,font=FS);d.rectangle(bb,fill='white');d.text((x,yy+7+j*31),t,font=FS,fill='black')
  else:
   d.rectangle((70,yy,W-70,yy+h),fill='white',outline='#333333',width=2)
   d.text((82,yy+6),'ref',font=FB,fill='black')
   for j,t in enumerate(lines):d.text((W/2+20,yy+10+j*31),t,font=FS,fill='black',anchor='mt')
 path=ROOT/f'diagram-{number:02}.png';im.save(path,dpi=(260,260));return path,im.size

def style_p(p,sz=13,bold=False):
 pf=p.paragraph_format;pf.left_indent=Inches(0);pf.right_indent=Inches(0);pf.first_line_indent=Inches(0)
 pf.space_before=Pt(0);pf.space_after=Pt(6);pf.line_spacing=1.08
 for run in p.runs:
  run.font.name='Times New Roman';run.font.size=Pt(sz);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
 return p

doc=Document(ROOT/'original.docx')
start=doc.paragraphs[785];anchor=doc.paragraphs[800]
for p in doc.paragraphs[786:800]:p._element.getparent().remove(p._element)
start.text='2.2.1 Các biểu đồ tuần tự (Sequence diagrams)';style_p(start,14,True)
start.paragraph_format.page_break_before=True
def p(text='',bold=False,size=13):
 p=doc.add_paragraph(text);style_p(p,size,bold);anchor._p.addprevious(p._p);return p
p('Các biểu đồ dưới đây mô tả thứ tự tương tác giữa tác nhân, giao diện, bộ phận xử lý nghiệp vụ và dữ liệu trong các trường hợp sử dụng của hệ thống bán quần áo. Phạm vi gồm các nhóm yêu cầu R1 đến R15 tại mục 1.5.1; các luồng xác thực được bổ sung theo chức năng đang triển khai.')
p('Quy ước đọc biểu đồ',True)
p('Thời gian đi từ trên xuống. Mũi tên nét liền biểu diễn yêu cầu; nét đứt biểu diễn kết quả trả về. Khung alt phân tách các nhánh thay thế, opt là nhánh tùy chọn, loop là lặp và ref tham chiếu một tương tác được mô tả riêng. Điều kiện đặt trong dấu ngoặc vuông.')
p('Các biểu đồ chức năng sử dụng mức trừu tượng phân tích: một đường sống xử lý có thể đại diện cho Controller và Service tương ứng; đường sống dữ liệu gộp AppDbContext, EF Core và SQL Server. Đây không phải lời gọi trực tiếp từ giao diện đến cơ sở dữ liệu. Sơ đồ chung thể hiện các thành phần trung gian; những sơ đồ sau lược bỏ phần trung gian lặp lại để tập trung vào nghiệp vụ.')
p('Các chức năng của khách hàng đã đăng nhập sử dụng userId lấy từ phiên xác thực. Các chức năng quản trị yêu cầu quyền Admin trên máy chủ. Lỗi dữ liệu, lỗi quyền truy cập và ngoại lệ được trả theo cơ chế chung; nếu nhánh lỗi kết thúc yêu cầu thì các bước thành công không được thực hiện.')
p('Danh mục các biểu đồ',True)
for lo,hi,label in [(1,7,'Luồng chung và xác thực'),(8,17,'Mua sắm và tài khoản khách hàng'),(18,27,'Quản trị cửa hàng')]:
 p(f'{label}: Hình 2.2.1-{lo:02} đến Hình 2.2.1-{hi:02}.')
p('Tài liệu ký pháp: OMG Unified Modeling Language 2.5.1, Chương 17 Interactions (https://www.omg.org/spec/UML/2.5.1/PDF).',size=11)
manifest=[]
for i,s in enumerate(specs,1):
 head=p(f'{i}. {s["title"]}',True,14);head.paragraph_format.page_break_before=True;head.paragraph_format.keep_with_next=True
 q=p(f'Yêu cầu liên quan: {s["req"]}. Điều kiện: {s["pre"]}',size=12);q.paragraph_format.keep_with_next=True
 path,size=render(s,i)
 q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True
 # Keep each figure and its explanation readable on one letter-size page.
 width=min(6.5,5.85*size[0]/size[1])
 q.add_run().add_picture(str(path),width=Inches(width))
 cap=p(f'Hình 2.2.1-{i:02}. Sơ đồ tuần tự {s["title"].lower()}',size=12);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=True
 p(s['desc'],size=12)
 manifest.append(dict(index=i,title=s['title'],image=str(path),width=width,size=size))
anchor.paragraph_format.page_break_before=True
doc.save(OUT)
(ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/'specs.json').write_text(json.dumps(specs,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Created {OUT}; {len(specs)} diagrams')
